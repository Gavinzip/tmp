#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path
import subprocess
import tempfile

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS_DIR = ROOT / "results"
FIGURES_DIR = ROOT / "figures"
REPORT_DIR = ROOT / "report"
OUT = REPORT_DIR / "term_project2_report_formatted.docx"

FORMULA_DIR = FIGURES_DIR / "formulas"
SYSTEM_PYTHON = Path("/Library/Frameworks/Python.framework/Versions/3.13/bin/python3")

BODY_FONT = "Arial Unicode MS"


def pct(value: float | int | None, digits: int = 2) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value) * 100:.{digits}f}%"


def money(value: float | int | None) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):,.0f}"


def set_run_font(run, size: int | None = None, bold: bool | None = None) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def paragraph(doc: Document, text: str = "", align: int | None = None, size: int = 11, bold: bool = False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    run = p.add_run(text)
    set_run_font(run, size=13 if level == 1 else 11, bold=True)
    run.font.color.rgb = RGBColor(0, 0, 0)


def bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_run_font(run, size=10)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False, size: int = 9, align: int | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "F2F2F2")
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def add_image(doc: Document, path: Path, caption: str, width: float = 6.15) -> None:
    if not path.exists():
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(caption)
    set_run_font(run, size=9)
    run.italic = True


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    styles = doc.styles
    for style_name in ["Normal", "List Bullet", "List Bullet 2"]:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(11 if style_name == "Normal" else 10)
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.line_spacing = 1.15


MATH_RENDER_SCRIPT = r"""
import json
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def render_math_png(formula, out_path, fontsize=28):
    fig = plt.figure(figsize=(13.0, 0.90), dpi=240)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    ax.text(0.0, 0.50, f"${formula}$", fontsize=fontsize, ha="left", va="center", color="black")
    fig.savefig(out_path, transparent=True, bbox_inches="tight", pad_inches=0.025)
    plt.close(fig)


def main():
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = Path(spec["path"])
    out.parent.mkdir(parents=True, exist_ok=True)
    render_math_png(spec["formula"], out, fontsize=int(spec.get("fontsize", 30)))


if __name__ == "__main__":
    main()
"""


def render_formula_image(path: Path, formula: str, fontsize: int = 30) -> Path:
    if SYSTEM_PYTHON.exists():
        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            spec_path = tmp_dir / "formula_spec.json"
            script_path = tmp_dir / "render_formula.py"
            spec_path.write_text(
                json.dumps(
                    {
                        "path": str(path),
                        "formula": formula,
                        "fontsize": fontsize,
                    },
                ),
                encoding="utf-8",
            )
            script_path.write_text(MATH_RENDER_SCRIPT, encoding="utf-8")
            result = subprocess.run(
                [str(SYSTEM_PYTHON), str(script_path), str(spec_path)],
                text=True,
                capture_output=True,
            )
            if result.returncode == 0 and path.exists():
                return path
            print(f"Formula image math renderer failed for {path}.")
            if result.stderr:
                print(result.stderr)
    raise RuntimeError(f"Unable to render formula image: {path}")


FormulaRow = tuple[str, str, Path]


def build_formula_images(prefix: str, rows: list[tuple[str, str]], fontsize: int = 30) -> list[FormulaRow]:
    FORMULA_DIR.mkdir(parents=True, exist_ok=True)
    out: list[FormulaRow] = []
    for idx, (label, formula) in enumerate(rows, start=1):
        path = FORMULA_DIR / f"{prefix}_{idx:02d}.png"
        render_formula_image(path, formula, fontsize=fontsize)
        out.append((label, formula, path))
    return out


def build_problem1_formula_images() -> list[FormulaRow]:
    return build_formula_images(
        "problem1",
        [
            ("每檔參數", r"S_i,\ L_i,\ M_i"),
            ("AAPL 範例", r"S_{\mathrm{AAPL}}=0.20,\ L_{\mathrm{AAPL}}=75,\ M_{\mathrm{AAPL}}=99"),
            ("移動平均", r"MA_t^{(i)}=\frac{1}{L_i}\sum_{k=0}^{L_i-1}P_{t-k}^{(i)}"),
            ("動能", r"Mom_t^{(i)}=\frac{P_t^{(i)}}{P_{t-M_i}^{(i)}}-1"),
            ("持有高點", r"H_t^{(i)}=\max(P_{\mathrm{entry}}^{(i)},\ldots,P_t^{(i)})"),
            ("回撤", r"DD_t^{(i)}=\frac{P_t^{(i)}}{H_t^{(i)}}-1"),
            ("出場", r"x_t^{(i)}=0\quad \mathrm{if}\quad x_{t-1}^{(i)}=1,\ DD_t^{(i)}\leq -S_i,\ P_t^{(i)}<MA_t^{(i)}"),
            ("進場", r"x_t^{(i)}=1\quad \mathrm{if}\quad x_{t-1}^{(i)}=0,\ P_t^{(i)}>MA_t^{(i)},\ Mom_t^{(i)}>0"),
            ("其他", r"x_t^{(i)}=x_{t-1}^{(i)}"),
            ("資產更新", r"V_t=V_{t-1}\left[1+x_{t-1}\left(\frac{P_t}{P_{t-1}}-1\right)-0.001|x_t-x_{t-1}|\right]"),
        ],
        fontsize=28,
    )


def build_problem2_formula_images() -> list[FormulaRow]:
    return build_formula_images(
        "problem2",
        [
            ("價格轉換", r"X_t=\ln(P_t^V),\quad Y_t=\ln(P_t^{MA})"),
            ("配對迴歸", r"X_t=\alpha+\beta Y_t+\varepsilon_t"),
            ("價差", r"Spread_t=X_t-\alpha-\beta Y_t"),
            ("標準化分數", r"z_t=\frac{Spread_t-\mu_L(Spread)}{\sigma_L(Spread)}"),
            ("Fixed 參數", r"L=40,\ entry_z=1.5,\ exit_z=0,\ stop_z=3.5,\ hold=60"),
            ("Long spread", r"z_t\leq -entry_z\Rightarrow +V-\beta MA"),
            ("Short spread", r"z_t\geq entry_z\Rightarrow -V+\beta MA"),
            ("均值回歸出場", r"Long:\ z_t\geq -exit_z,\quad Short:\ z_t\leq exit_z"),
            ("風控", r"|z_t|\geq stop_z\quad \mathrm{or}\quad holding\ days\geq max\ hold"),
            ("GA 搜尋", r"chromosome=(L,\ entry_z,\ exit_z,\ stop_z,\ holding\ days)"),
        ],
        fontsize=28,
    )


def add_formula_list(doc: Document, rows: list[FormulaRow]) -> None:
    for label, formula, path in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        run = p.add_run(label)
        set_run_font(run, size=10, bold=True)

        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(5)
        p_img.paragraph_format.keep_together = True
        p_img.add_run().add_picture(str(path), width=Inches(formula_image_width(path, formula)))


def formula_image_width(path: Path, formula: str) -> float:
    """Keep short formulas compact while capping long formulas to the text width."""
    try:
        with Image.open(path) as image:
            natural_width = image.size[0] / 240 * 0.82
    except OSError:
        natural_width = len(formula) * 0.06
    return max(1.35, min(5.85, natural_width))


def load_data() -> dict[str, pd.DataFrame]:
    return {
        "p1_summary": pd.read_csv(RESULTS_DIR / "reco_problem1_per_stock_summary.csv"),
        "p1_params": pd.read_csv(RESULTS_DIR / "reco_problem1_ga_params.csv"),
        "p1_turns": pd.read_csv(RESULTS_DIR / "reco_problem1_turning_points.csv"),
        "p1_phases": pd.read_csv(RESULTS_DIR / "reco_problem1_stock_phases.csv"),
        "p2_metrics": pd.read_csv(RESULTS_DIR / "reco_problem2_ga_window_metrics.csv"),
        "p2_params": pd.read_csv(RESULTS_DIR / "reco_problem2_ga_params.csv"),
        "p2_trades": pd.read_csv(RESULTS_DIR / "reco_problem2_ga_trade_log.csv"),
    }


def strategy_return(summary: pd.DataFrame, ticker: str, strategy: str) -> float:
    row = summary[(summary["ticker"] == ticker) & (summary["strategy"] == strategy)]
    return float(row.iloc[0]["cumulative_return"]) if not row.empty else float("nan")


def p1_parameter_rows(data: dict[str, pd.DataFrame]) -> list[list[object]]:
    rows = []
    summary = data["p1_summary"]
    params = data["p1_params"]
    for _, r in params.iterrows():
        ticker = r["ticker"]
        ga_ret = strategy_return(summary, ticker, "ga_trend_stop_momentum")
        fixed_ret = strategy_return(summary, ticker, "fixed_trend_stop_momentum")
        dca_ret = strategy_return(summary, ticker, "dca")
        bh_ret = strategy_return(summary, ticker, "buy_and_hold")
        if ga_ret > fixed_ret and ga_ret > dca_ret:
            note = "GA 對進出場有改善"
        elif ga_ret > dca_ret:
            note = "優於 DCA，但未必贏 Fixed"
        else:
            note = "防守有效，但報酬不突出"
        rows.append(
            [
                ticker,
                f"停損 {float(r['ga_trailing_stop_pct']) * 100:.1f}%, 均線 {int(r['ga_reentry_ma'])}, 動能 {int(r['ga_momentum_lookback'])}",
                pct(ga_ret),
                pct(fixed_ret),
                pct(dca_ret),
                pct(bh_ret),
                note,
            ]
        )
    return rows


def selected_stock_signals(turns: pd.DataFrame, ticker: str) -> list[pd.Series]:
    g = turns[turns["ticker"] == ticker].copy()
    selected: list[pd.Series] = []
    exits = g[g["action"] == "Exit"]
    reentries = g[g["action"] == "Re-entry"]
    if not exits.empty:
        selected.append(exits.reindex(exits["drawdown_from_peak"].abs().sort_values(ascending=False).index).iloc[0])
    if not reentries.empty:
        selected.append(reentries.reindex(reentries["price_return_until_next_change"].abs().sort_values(ascending=False).index).iloc[0])
    latest = g.sort_values("date").tail(1)
    if not latest.empty and not any(str(s["date"]) == str(latest.iloc[0]["date"]) for s in selected):
        selected.append(latest.iloc[0])
    return selected[:3]


def signal_sentence(row: pd.Series) -> str:
    action = "出場" if row["action"] == "Exit" else "重新進場"
    next_change = pct(row["price_return_until_next_change"])
    if row["action"] == "Exit":
        reason = "當時回撤已達停損條件，且價格低於 GA 選出的均線，因此策略先降為空手。"
    else:
        reason = "當時價格重新站上 GA 選出的均線且動能轉正，因此策略恢復持有。"
    return (
        f"{row['date']} {action}：回撤 {pct(row['drawdown_from_peak'])}、"
        f"63 日波動率 {pct(row['volatility_63d'])}。{reason}"
        f"到下一次訊號前，股價變化約 {next_change}。"
    )


def stock_summary_text(data: dict[str, pd.DataFrame], ticker: str) -> str:
    summary = data["p1_summary"]
    params = data["p1_params"]
    phases = data["p1_phases"]
    p = params[params["ticker"] == ticker].iloc[0]
    ph = phases[phases["ticker"] == ticker].iloc[0]
    ga_ret = strategy_return(summary, ticker, "ga_trend_stop_momentum")
    fixed_ret = strategy_return(summary, ticker, "fixed_trend_stop_momentum")
    dca_ret = strategy_return(summary, ticker, "dca")
    bh_ret = strategy_return(summary, ticker, "buy_and_hold")
    return (
        f"{ticker} 的 GA 參數為停損 {float(p['ga_trailing_stop_pct']) * 100:.1f}%、"
        f"均線 {int(p['ga_reentry_ma'])} 日、動能 {int(p['ga_momentum_lookback'])} 日。"
        f"2016-2026 測試期中，GA 累積報酬 {pct(ga_ret)}，Fixed {pct(fixed_ret)}，"
        f"DCA {pct(dca_ret)}，Buy-and-Hold {pct(bh_ret)}。"
        f"該股測試期最大價格回撤約 {pct(ph['worst_price_drawdown'])}，"
        f"策略最大回撤約 {pct(ph['worst_strategy_drawdown'])}。"
    )


def p2_metric_rows(metrics: pd.DataFrame, window: str = "2016-2026") -> list[list[object]]:
    labels = {
        "buy_and_hold_pair": "Buy-and-Hold",
        "dca_pair": "DCA",
        "fixed_pairs_zscore": "Classic V-MA Z-score",
        "ga_pairs_zscore": "GA V-MA Z-score",
    }
    order = ["buy_and_hold_pair", "dca_pair", "fixed_pairs_zscore", "ga_pairs_zscore"]
    w = metrics[metrics["window"].astype(str) == window]
    rows = []
    for strategy in order:
        r = w[w["strategy"] == strategy]
        if r.empty:
            continue
        r = r.iloc[0]
        rows.append(
            [
                labels[strategy],
                pct(r["cumulative_return"]),
                pct(r["max_drawdown"]),
                pct(r["volatility"]),
                money(r["end_equity"]),
            ]
        )
    return rows


def p2_parameter_rows(params: pd.DataFrame) -> list[list[object]]:
    rows: list[list[object]] = [
        ["Classic Fixed", "2016-2026", "V-MA", 40, "1.50", "0.00", "3.50", 60],
    ]
    p = params[params["window"].astype(str) == "2016-2026"]
    if not p.empty:
        r = p.iloc[0]
        rows.append(
            [
                "GA",
                "2016-2026",
                r["selected_pairs"],
                int(r["lookback"]),
                f"{float(r['entry_z']):.2f}",
                f"{float(r['exit_z']):.2f}",
                f"{float(r['stop_z']):.2f}",
                int(r["max_holding_days"]),
            ]
        )
    return rows


def p2_result_sentence(metrics: pd.DataFrame) -> str:
    w = metrics[metrics["window"].astype(str) == "2016-2026"].set_index("strategy")
    ga = float(w.loc["ga_pairs_zscore", "cumulative_return"]) if "ga_pairs_zscore" in w.index else float("nan")
    fixed = float(w.loc["fixed_pairs_zscore", "cumulative_return"]) if "fixed_pairs_zscore" in w.index else float("nan")
    bh = float(w.loc["buy_and_hold_pair", "cumulative_return"]) if "buy_and_hold_pair" in w.index else float("nan")
    dca = float(w.loc["dca_pair", "cumulative_return"]) if "dca_pair" in w.index else float("nan")
    return (
        f"2016-2026 正式測試窗中，GA 累積報酬 {pct(ga)}，Classic Fixed {pct(fixed)}，"
        f"Buy-and-Hold {pct(bh)}，DCA {pct(dca)}。"
        "GA 比 Classic Fixed 少虧，但兩者仍明顯輸給長期持有型基準。"
    )


def p2_trade_bullets(doc: Document, trades: pd.DataFrame) -> None:
    selected = []
    for strategy in ["fixed_pairs_zscore", "ga_pairs_zscore"]:
        s = trades[(trades["window"].astype(str) == "2016-2026") & (trades["strategy"] == strategy)].copy()
        if s.empty:
            continue
        s["abs_ret"] = s["trade_return"].abs()
        selected.append(s.sort_values("abs_ret", ascending=False).head(2))
    if not selected:
        return
    picked = pd.concat(selected).head(4)
    for _, r in picked.iterrows():
        side = "做多價差" if r["side"] == "long_spread" else "做空價差"
        reverted = (r["side"] == "long_spread" and r["exit_z"] > r["entry_z"]) or (r["side"] == "short_spread" and r["exit_z"] < r["entry_z"])
        result = "z-score 有往 0 靠近" if reverted else "z-score 沒有順利回到 0"
        strategy_label = "Classic V-MA Z-score" if r["strategy"] == "fixed_pairs_zscore" else "GA V-MA Z-score"
        reason_label = {
            "mean_reversion": "均值回歸",
            "max_holding_days": "達到最大持有天數",
            "stop_z": "價差停損",
            "open_at_end": "測試期結束仍持倉",
        }.get(str(r["exit_reason"]), str(r["exit_reason"]))
        bullet(
            doc,
            f"{r['entry_date']} 至 {r['exit_date']}，{strategy_label}，"
            f"{side}，z-score {float(r['entry_z']):.2f} -> {float(r['exit_z']):.2f}，"
            f"單筆報酬 {pct(r['trade_return'])}。{result}，出場原因為 {reason_label}。",
        )


def add_title_page(doc: Document) -> None:
    paragraph(doc, "Term Project II:", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    paragraph(doc, "量化交易策略設計與 GA 參數搜尋", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    paragraph(doc, "Problem 1：個股規則型策略\nProblem 2：Classic V-MA Z-score Pairs Trading", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    doc.add_paragraph()


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    problem1_formulas = build_problem1_formula_images()
    problem2_formulas = build_problem2_formula_images()
    data = load_data()

    doc = Document()
    setup_document(doc)
    add_title_page(doc)

    heading(doc, "一、背景與目的", 1)
    paragraph(
        doc,
        "本專題針對 AAPL、MSFT、V、MA、JPM 五檔股票設計可解釋的規則型交易策略，"
        "並使用 2016 至 2026 年的樣本外測試期比較策略表現。"
        "研究重點不是預測明天漲跌，而是讓每一個交易決策都能由明確規則、參數與歷史資料驗證來說明。",
    )
    paragraph(
        doc,
        "Problem 1 著重在每一檔股票各自的進出場規則；Problem 2 則固定使用 Visa 與 Mastercard "
        "這組同產業支付網路 pair，測試 Classic V-MA Z-score Pairs Trading 與 GA 參數搜尋的差異。",
    )

    heading(doc, "二、資料與驗證方式", 1)
    bullet(doc, "資料來源為 Yahoo Finance 調整後收盤價，主要測試期間為 2016-2026。")
    bullet(doc, "所有 GA 參數只使用測試期以前的資料訓練，不使用測試期結果回頭調參。")
    bullet(doc, "比較基準包含 Buy-and-Hold 與 DCA，避免只看策略本身而無法判斷相對表現。")
    bullet(doc, "績效指標包含累積報酬、最大回撤、波動率與期末資產。")

    heading(doc, "三、Problem 1：個股策略設計", 1)
    paragraph(
        doc,
        "Problem 1 不把五檔股票合成一條線，而是每檔股票各自套用同一套規則。"
        "GA 只負責替每檔股票挑選停損、均線視窗與動能視窗，實際交易仍由固定規則決定。",
    )
    paragraph(doc, "Problem 1 公式如下。項目名稱與說明使用 Word 原生文字排版，公式本身各自以單張 PNG 插入。", size=10)
    add_formula_list(doc, problem1_formulas)

    doc.add_page_break()
    heading(doc, "3.1 個股參數與績效摘要", 2)
    add_table(
        doc,
        ["Stock", "GA Parameters", "GA", "Fixed", "DCA", "Buy-and-Hold", "Interpretation"],
        p1_parameter_rows(data),
        [0.55, 1.65, 0.65, 0.65, 0.65, 0.85, 1.45],
    )

    for idx, ticker in enumerate(["AAPL", "MSFT", "V", "MA", "JPM"], start=2):
        heading(doc, f"3.{idx} {ticker} 個股結果與關鍵訊號", 2)
        paragraph(doc, stock_summary_text(data, ticker))
        add_image(doc, FIGURES_DIR / f"reco_problem1_{ticker}_vs_benchmarks.png", f"圖：{ticker} 個股策略與基準比較。", width=6.15)
        bullet(doc, "關鍵進出場點：")
        for sig in selected_stock_signals(data["p1_turns"], ticker):
            bullet(doc, signal_sentence(sig), level=1)

    doc.add_page_break()
    heading(doc, "四、Problem 2：Classic V-MA Z-score Pairs Trading", 1)
    paragraph(
        doc,
        "Problem 2 固定使用 Visa 與 Mastercard，不在測試期之後回頭更換 pair。"
        "Classic 版本使用固定參數；GA 版本只調整交易參數，不改變 pair 本身。",
    )
    paragraph(doc, "Problem 2 公式如下。文字說明維持 Word 原生排版，公式本身各自以單張 PNG 插入。", size=10)
    add_formula_list(doc, problem2_formulas)

    heading(doc, "4.1 2016-2026 主要結果", 2)
    add_table(
        doc,
        ["Strategy", "Cumulative", "Max Drawdown", "Volatility", "End Equity"],
        p2_metric_rows(data["p2_metrics"], "2016-2026"),
        [1.45, 0.85, 0.85, 0.85, 1.0],
    )
    paragraph(
        doc,
        "2016-2026 長期結果顯示，Buy-and-Hold 與 DCA 因為長期持有 V-MA 支付網路龍頭，報酬明顯高於配對交易。"
        "配對交易的目的不是追求單邊多頭報酬，而是測試價差偏離後是否能透過均值回歸取得較低波動的策略表現。",
    )

    add_image(doc, FIGURES_DIR / "reco_problem2_ga_vs_fixed.png", "圖：Problem 2 四種方法比較，包含 Buy-and-Hold、DCA、Classic 與 GA。", width=6.15)
    add_image(doc, FIGURES_DIR / "reco_problem2_2016-2026_v_ma_price_position.png", "圖：V-MA 價格、資產曲線、z-score 與持倉變化。", width=6.15)

    heading(doc, "4.2 GA Parameters 與關鍵交易", 2)
    add_table(
        doc,
        ["Strategy", "Window", "Pair", "lookback", "entry_z", "exit_z", "stop_z", "max hold"],
        p2_parameter_rows(data["p2_params"]),
        [1.15, 0.9, 0.75, 0.65, 0.65, 0.65, 0.65, 0.75],
    )
    bullet(doc, p2_result_sentence(data["p2_metrics"]))
    bullet(doc, "Interpretation：V 與 MA 在 2016-2026 同向長期上漲，Pair Trading 的 market-neutral 結構會抵消大部分單邊趨勢，所以不適合拿來追求 Buy-and-Hold 型報酬。")
    bullet(doc, "GA Interpretation：GA 只用 2015 年底以前資料挑參數，因此它能降低虧損與回撤，但無法保證測試期一定轉正。")
    bullet(doc, "關鍵交易：")
    p2_trade_bullets(doc, data["p2_trades"])

    heading(doc, "五、結論", 1)
    bullet(doc, "Problem 1 的重點是個股差異：同一套規則放到不同股票，會因趨勢強弱、回撤速度與波動型態而有不同效果。")
    bullet(doc, "GA 可以提供一組事前選出的參數，但不代表一定打敗 Buy-and-Hold；強多頭股票常會讓長抱策略勝出。")
    bullet(doc, "Problem 2 的 V-MA pair trading 已具備進場、出場、停損、最大持有天數與持倉變化，不是沒有交易的水平線。")
    bullet(doc, "本報告保守解釋 GA：GA 是參數搜尋工具，不是保證獲利模型；結果必須用時間序列驗證檢查。")

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
